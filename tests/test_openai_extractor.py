from datetime import date
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
from typing import Any, cast

from openai import OpenAI
from openpyxl import Workbook
import pytest

from tt_automation.config import Settings
from tt_automation.extraction.documents import SourceDocument
from tt_automation.extraction.openai_extractor import (
    ExtractionError,
    extract_transfer_data,
)
from tt_automation.models import TransferData


class FakeResponses:
    def __init__(self, parsed: TransferData) -> None:
        self.parsed = parsed
        self.request: dict[str, Any] = {}

    def parse(self, **kwargs: Any) -> SimpleNamespace:
        self.request = kwargs
        return SimpleNamespace(output_parsed=self.parsed)


class FakeOpenAI:
    def __init__(self, parsed: TransferData) -> None:
        self.responses = FakeResponses(parsed)


class InvalidResponses:
    def parse(self, **kwargs: Any) -> SimpleNamespace:
        TransferData.model_validate({"invoice_date": "not a date"})
        raise AssertionError("Expected validation to fail")


class InvalidOpenAI:
    def __init__(self) -> None:
        self.responses = InvalidResponses()


def test_extracts_typed_data_from_image_and_workbook_without_network() -> None:
    workbook = Workbook()
    workbook.active["A1"] = "Invoice INV-42"
    workbook_buffer = BytesIO()
    workbook.save(workbook_buffer)
    expected = TransferData(
        beneficiary_name="Example Ltd",
        currency="USD",
        amount=Decimal("125.00"),
        invoice_number="INV-42",
        invoice_date=date(2026, 7, 3),
    )
    fake_client = FakeOpenAI(expected)
    documents = [
        SourceDocument("invoice.png", b"fake image bytes"),
        SourceDocument("details.xlsx", workbook_buffer.getvalue()),
    ]

    result = extract_transfer_data(
        documents,
        Settings(
            openai_api_key="test",
            openai_model="test-model",
            openai_base_url="https://example.test/openai/v1/",
        ),
        client=cast(OpenAI, fake_client),
    )

    assert result == expected
    assert fake_client.responses.request["model"] == "test-model"
    assert fake_client.responses.request["reasoning"] == {"effort": "low"}
    request_text = repr(fake_client.responses.request["input"])
    assert "data:image/png;base64," in request_text
    assert "Workbook source: details.xlsx" in request_text
    assert "Invoice INV-42" in request_text


def test_converts_invalid_structured_output_to_extraction_error() -> None:
    settings = Settings(
        openai_api_key="test",
        openai_model="test-model",
        openai_base_url="https://example.test/openai/v1/",
    )

    with pytest.raises(ExtractionError, match="unsupported format"):
        extract_transfer_data(
            [SourceDocument("invoice.png", b"fake image bytes")],
            settings,
            client=cast(OpenAI, InvalidOpenAI()),
        )
