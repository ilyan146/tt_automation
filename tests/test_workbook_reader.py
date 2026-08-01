from datetime import date
from io import BytesIO

from openpyxl import Workbook

from tt_automation.config import TEMPLATE_PATH
from tt_automation.extraction.documents import SourceDocument
from tt_automation.extraction.workbook_reader import workbook_to_text


def test_reads_modern_workbook_with_cell_positions() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Invoice"
    sheet["B2"] = "Invoice number"
    sheet["C2"] = "INV-42"
    sheet["C3"] = date(2026, 7, 3)
    buffer = BytesIO()
    workbook.save(buffer)

    text = workbook_to_text(SourceDocument("invoice.xlsx", buffer.getvalue()))

    assert "[Sheet: Invoice]" in text
    assert 'Row 2: B="Invoice number" | C="INV-42"' in text
    assert 'Row 3: C="2026-07-03T00:00:00"' in text


def test_reads_supplied_legacy_template() -> None:
    document = SourceDocument(TEMPLATE_PATH.name, TEMPLATE_PATH.read_bytes())

    text = workbook_to_text(document)

    assert "[Sheet: Sheet1 (2)]" in text
    assert 'D="26ME230"' in text
    assert 'L="BKCHCNBJ550"' in text
